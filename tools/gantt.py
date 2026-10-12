"""Generate reports/gantt.xlsx from PLAN.md and git history.

PLAN.md is the source of truth. Planned dates are spread between the
header's Inicio/Fin proportionally to each item's `est:` (days), phases
sequential. Real dates come from the commit that marked an item [x]
(first commit whose diff turns the item's checkbox to [x]). Items with
a scope decision (DESCARTADO / POSTPUESTO / DESEABLE A FUTURO) leave the
progress denominator. Progress and deviation are Excel formulas.
"""

from __future__ import annotations

import re
import subprocess
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PLAN = ROOT / "PLAN.md"
OUT = ROOT / "reports" / "gantt.xlsx"

RE_PHASE = re.compile(r"^### Fase (\d+) — (.+?)\s+\(est: ([\d.,]+) d[ií]as?\)")
RE_ITEM = re.compile(r"^- \[([ x])\] (\d+\.\d+) (.+)$")
RE_EST = re.compile(r"\(est: ([\d.,]+)\)")
RE_DATE = re.compile(r"^(Inicio|Fin): (\d{4}-\d{2}-\d{2})")
SCOPE_STATES = {
    "DESCARTADO": "Descartado",
    "POSTPUESTO": "Postpuesto",
    "DESEABLE A FUTURO": "Deseable a futuro",
}


def num(s: str) -> float:
    return float(s.replace(",", "."))


def parse_plan():
    text = PLAN.read_text(encoding="utf-8")
    lines = text.splitlines()
    start = end = None
    for ln in lines[:8]:
        for part in ln.split(" · "):
            m = RE_DATE.match(part.strip())
            if m:
                d = date.fromisoformat(m.group(2))
                if m.group(1) == "Inicio":
                    start = d
                else:
                    end = d
    if not (start and end):
        raise SystemExit("PLAN.md sin Inicio:/Fin: en el encabezado")

    phases = []
    cur = None
    i = 0
    while i < len(lines):
        ln = lines[i]
        m = RE_PHASE.match(ln)
        if m:
            cur = {
                "id": int(m.group(1)),
                "name": m.group(2),
                "est": num(m.group(3)),
                "items": [],
                "scope": None,
            }
            # phase-level scope note: lines until first item
            j = i + 1
            while (
                j < len(lines)
                and not RE_ITEM.match(lines[j])
                and not lines[j].startswith("###")
            ):
                for k, v in SCOPE_STATES.items():
                    if k in lines[j]:
                        cur["scope"] = v
                j += 1
            phases.append(cur)
            i += 1
            continue
        if ln.startswith(("## ", "### Calendario")):
            cur = None
        m = RE_ITEM.match(ln)
        if m and cur is not None:
            done = m.group(1) == "x"
            iid = m.group(2)
            # gather the item block (continuation lines)
            block = [m.group(3)]
            j = i + 1
            while j < len(lines) and lines[j].startswith("      "):
                block.append(lines[j].strip())
                j += 1
            full = " ".join(block)
            est_m = RE_EST.search(full)
            est = num(est_m.group(1)) if est_m else 0.0
            title = full.split("Aceptación:")[0]
            title = re.sub(r"\s*\(est: [\d.,]+\)", "", title).strip().rstrip(".")
            title = re.sub(r"\s+", " ", title)
            if len(title) > 110:
                title = title[:107].rstrip() + "..."
            scope = cur["scope"]
            for k, v in SCOPE_STATES.items():
                if k in full:
                    scope = v
            cur["items"].append(
                {"id": iid, "title": title, "est": est, "done": done, "scope": scope}
            )
        i += 1
    return start, end, phases


def real_dates():
    """Map item id -> date of the commit that first marked it [x]."""
    out = {}
    try:
        log = subprocess.run(
            [
                "git",
                "log",
                "--reverse",
                "--date=short",
                "--format=COMMIT %ad",
                "-p",
                "--",
                "PLAN.md",
            ],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=True,
        ).stdout
    except subprocess.CalledProcessError:
        return out
    cur = None
    for ln in log.splitlines():
        if ln.startswith("COMMIT "):
            cur = date.fromisoformat(ln.split()[1])
        elif ln.startswith("+- [x] ") and cur:
            m = re.match(r"\+- \[x\] (\d+\.\d+) ", ln)
            if m and m.group(1) not in out:
                out[m.group(1)] = cur
    return out


def build():
    start, end, phases = parse_plan()
    reals = real_dates()
    active = [p for p in phases if p["scope"] is None]
    total_est = (
        sum(it["est"] for p in active for it in p["items"] if it["scope"] is None) or 1
    )
    span = (end - start).days
    cursor = start

    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    phase_fill = PatternFill("solid", fgColor="D9E1F2")
    font = Font(name="Calibri", size=10)
    bold = Font(name="Calibri", size=10, bold=True)
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    ws["A1"] = "Bitácora de Viaje — Cronograma"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = (
        f"Inicio {start.isoformat()} · Fin {end.isoformat()} · generado {date.today().isoformat()} desde PLAN.md"
    )
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="595959")

    cols = [
        "ID",
        "Fase / Ítem",
        "Est. (días)",
        "Inicio plan",
        "Fin plan",
        "Inicio real",
        "Fin real",
        "Estado",
        "Avance",
        "Desvío (días)",
    ]
    HR = 4
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=HR, column=c, value=name)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = border

    # week columns
    week0 = start - timedelta(days=start.weekday())
    nweeks = (end - week0).days // 7 + 1
    GC = len(cols) + 1
    for w in range(nweeks):
        ws_date = week0 + timedelta(days=7 * w)
        cell = ws.cell(row=HR, column=GC + w, value=ws_date)
        cell.number_format = "dd-mmm"
        cell.font = Font(name="Calibri", size=8, bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", text_rotation=90
        )
        ws.column_dimensions[get_column_letter(GC + w)].width = 2.6
    ws.row_dimensions[HR].height = 48

    r = HR + 1
    phase_rows = []
    item_rows_all = []
    for p in phases:
        prow = r
        ws.cell(row=r, column=1, value=f"F{p['id']}")
        ws.cell(row=r, column=2, value=f"Fase {p['id']} — {p['name']}")
        ws.cell(row=r, column=3, value=p["est"])
        r += 1
        first = r
        for it in p["items"]:
            ws.cell(row=r, column=1, value=it["id"])
            ws.cell(row=r, column=2, value=it["title"])
            ws.cell(row=r, column=3, value=it["est"])
            if p["scope"] is None and it["scope"] is None:
                days = span * it["est"] / total_est
                i0 = cursor
                i1 = cursor + timedelta(days=max(days, 0.5))
                cursor = i1
                ws.cell(row=r, column=4, value=i0)
                ws.cell(row=r, column=5, value=i1)
            if it["id"] in reals:
                ws.cell(row=r, column=7, value=reals[it["id"]])
            if it["scope"]:
                state = it["scope"]
            elif it["done"]:
                state = "Completado"
            else:
                state = "Pendiente"
            ws.cell(row=r, column=8, value=state)
            ws.cell(
                row=r,
                column=9,
                value=1 if it["done"] else 0 if not it["scope"] else None,
            )
            ws.cell(row=r, column=10, value=f'=IF(AND(G{r}<>"",E{r}<>""),G{r}-E{r},"")')
            item_rows_all.append(r)
            r += 1
        last = r - 1
        if last >= first:
            ws.cell(
                row=prow,
                column=4,
                value=f'=IF(COUNT(D{first}:D{last})=0,"",MIN(D{first}:D{last}))',
            )
            ws.cell(
                row=prow,
                column=5,
                value=f'=IF(COUNT(E{first}:E{last})=0,"",MAX(E{first}:E{last}))',
            )
            ws.cell(
                row=prow,
                column=6,
                value=f'=IF(COUNT(F{first}:F{last})=0,"",MIN(F{first}:F{last}))',
            )
            ws.cell(
                row=prow,
                column=7,
                value=f'=IF(COUNT(G{first}:G{last})=0,"",MAX(G{first}:G{last}))',
            )
            ws.cell(
                row=prow,
                column=9,
                value=f'=IF(COUNT(I{first}:I{last})=0,"",AVERAGE(I{first}:I{last}))',
            )
            ws.cell(
                row=prow,
                column=10,
                value=f'=IF(AND(G{prow}<>"",E{prow}<>""),G{prow}-E{prow},"")',
            )
        ws.cell(
            row=prow,
            column=8,
            value=p["scope"]
            or f'=IF(I{prow}="","",IF(I{prow}=1,"Completada",IF(I{prow}=0,"Pendiente","En curso")))',
        )
        for c in range(1, GC + nweeks):
            ws.cell(row=prow, column=c).fill = phase_fill
            ws.cell(row=prow, column=c).font = bold
        phase_rows.append(prow)

    # global row
    r += 1
    ws.cell(row=r, column=2, value="AVANCE GLOBAL").font = bold
    rng = ",".join(f"I{x}" for x in item_rows_all)
    ws.cell(row=r, column=9, value=f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))').font = bold
    gl = r

    # formats
    for row in ws.iter_rows(min_row=HR + 1, max_row=gl, max_col=GC + nweeks - 1):
        for cell in row:
            if cell.font != bold:
                cell.font = font
            cell.border = border
            if cell.column in (4, 5, 6, 7):
                cell.number_format = "yyyy-mm-dd"
            if cell.column == 9:
                cell.number_format = "0%"
            if cell.column in (3, 10):
                cell.number_format = "0.0"
    widths = {1: 6, 2: 70, 3: 9, 4: 11, 5: 11, 6: 11, 7: 11, 8: 16, 9: 8, 10: 10}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=HR + 1, column=3)

    # Gantt bars: planned (blue) and real (green) via conditional formatting
    first_w = get_column_letter(GC)
    last_w = get_column_letter(GC + nweeks - 1)
    area = f"{first_w}{HR + 1}:{last_w}{gl - 1}"
    hdr = f"{first_w}${HR}"
    plan_rule = FormulaRule(
        formula=[
            f'AND($D{HR + 1}<>"",$E{HR + 1}<>"",{hdr}<=$E{HR + 1},{hdr}+6>=$D{HR + 1})'
        ],
        fill=PatternFill("solid", fgColor="5B9BD5"),
    )
    real_rule = FormulaRule(
        formula=[
            f'AND($G{HR + 1}<>"",{hdr}<=$G{HR + 1},{hdr}+6>=IF($F{HR + 1}<>"",$F{HR + 1},$D{HR + 1}))'
        ],
        fill=PatternFill("solid", fgColor="70AD47"),
    )
    ws.conditional_formatting.add(area, real_rule)
    ws.conditional_formatting.add(area, plan_rule)
    today_rule = FormulaRule(
        formula=[f"AND({hdr}<=TODAY(),{hdr}+6>=TODAY())"],
        fill=PatternFill("solid", fgColor="FFF2CC"),
    )
    ws.conditional_formatting.add(f"{first_w}{HR}:{last_w}{HR}", today_rule)

    # legend
    lr = gl + 2
    ws.cell(
        row=lr,
        column=2,
        value="Leyenda: azul = planificado · verde = real · amarillo = semana actual. "
        "Estados Descartado / Postpuesto / Deseable a futuro no cuentan para el avance.",
    ).font = Font(name="Calibri", size=9, italic=True, color="595959")

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    n_items = len(item_rows_all)
    n_active = sum(
        1 for p in phases for it in p["items"] if not it["scope"] and p["scope"] is None
    )
    print(
        f"{OUT.relative_to(ROOT)}: {len(phases)} fases, {n_items} ítems ({n_active} activos), "
        f"{nweeks} semanas, {len(reals)} con fecha real"
    )


if __name__ == "__main__":
    build()
